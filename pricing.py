"""قراءة شيت الأسعار من ملف xlsx إلى صفوف منظمة.

القواعد:
- الشيت المخفي (Hide sheet) = عرض قديم، البوت يتجاهله.
- العملة تنقرأ من تنسيق الخلية: إذا بيه $ فهو دولار، وإذا الرقم 10,000 وفوق بدون رمز فهو دينار.
- الخلايا المدموجة عمودياً تتكرر قيمتها على كل صفوفها (مثل اسم الفندق لأربع مدد).
- كل خلية لها نص للشركات (كامل) ونص للزبون (مقصوص منه أي ذكر للعمولة).
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

TATWEEL = "ـ"
MAX_CELL_CHARS = 1500
MIN_CELLS = 10  # شيت بيه أقل من هذا يعتبر فارغ

_COMMISSION = re.compile(r"عمول")
_TRAILING_JUNK = re.compile(r"[\s(\[\-–—:،,+*]+$")
_MERGE = re.compile(r'<mergeCell ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"')
_SPACED_NUMBER = re.compile(r"^\d{1,3}(?: \d{3})+$")
_INVISIBLE = re.compile("[\u200b-\u200f\u202a-\u202e\u2066-\u2069\ufeff" + TATWEEL + "]")


def clean(text) -> str:
    """سطر واحد: يشيل المدّات (ـ) والفراغات الزايدة."""
    return " ".join(_INVISIBLE.sub("", str(text)).split())


def clean_multiline(text) -> str:
    """يحافظ على الأسطر، ويحول خطوط الفصل (-----) إلى سطر جديد."""
    text = _INVISIBLE.sub("", str(text)).replace("\r", "\n")
    text = re.sub(r"[-*=_]{3,}", "\n", text)
    lines = [" ".join(line.split()) for line in text.split("\n")]
    return "\n".join(line for line in lines if line)[:MAX_CELL_CHARS]


_AMOUNT = r"[\d.,]+\s*(?:\$|دينار|د\.?\s?ع|الف|دولار)?"
_CLAUSE = re.compile(
    r"\(?\s*عمول[ةه]\s*(?:الشركات|شركات|العرض الخاص)?\s*" + _AMOUNT
    + r"(?:\s*(?:و|\+)?\s*(?:عمول[ةه]\s*)?العرض الخاص\s*(?:اسعار تسديد|تسديد|" + _AMOUNT + r"))?\s*\)?"
)


def strip_commission(text: str) -> str:
    """يشيل جملة العمولة من النص ويخلي الباقي (مثل عنوان العرض). يرجع نص فارغ إذا الخلية كلها عمولة."""
    if not _COMMISSION.search(text):
        return text
    text = _CLAUSE.sub(" ", text)
    match = _COMMISSION.search(text)
    if match:  # صيغة ما نعرفها: نقص من أول ذكر للعمولة لآخر الخلية
        text = text[: match.start()]
    text = re.sub(r"\(\s*\)", " ", text)
    lines = [_TRAILING_JUNK.sub("", " ".join(line.split())) for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


@dataclass
class Cell:
    text: str  # النص الكامل (للشركات)
    number: float | None = None  # القيمة إذا الخلية رقم
    unit: str = ""  # "$" أو "د.ع" أو فارغ إذا غير معروف
    original: bool = True  # False إذا القيمة مكررة من خلية مدموجة فوقها
    span_to: int = 0  # آخر عمود إذا الخلية مدموجة أفقياً

    @property
    def customer_text(self) -> str:
        return strip_commission(self.text)

    @property
    def has_commission(self) -> bool:
        return bool(_COMMISSION.search(self.text))


def _number_cell(value: float, number_format: str) -> Cell:
    text = f"{value:,.0f}" if float(value).is_integer() else f"{value:,.2f}"
    nf = number_format or ""
    if "$" in nf or "USD" in nf:
        unit = "$"
    elif "IQD" in nf or "د.ع" in nf or "دينار" in nf or abs(value) >= 10000:
        unit = "د.ع"
    else:
        unit = ""  # رقم صغير بدون رمز: الريندر يحدد معناه من عنوان العمود
    return Cell(text=text, number=float(value), unit=unit)


def _make_cell(raw) -> Cell | None:
    value = raw.value
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return _number_cell(value, getattr(raw, "number_format", "") or "")
    if isinstance(value, (datetime, date)):
        return Cell(text=value.strftime("%d-%m-%Y"))
    text = clean_multiline(value)
    if not text:
        return None
    if _SPACED_NUMBER.match(text):  # مثل "50 000" مكتوبة كنص
        return _number_cell(float(text.replace(" ", "")), "")
    return Cell(text=text)


@dataclass
class Sheet:
    key: str  # اسم الشيت الأصلي
    title: str  # الاسم بعد التنظيف
    rows: list[tuple[int, dict[int, Cell]]] = field(default_factory=list)  # (رقم الصف, {رقم العمود: خلية})
    cells: int = 0

    @property
    def commission_notes(self) -> list[str]:
        seen: list[str] = []
        for _, cells in self.rows:
            for cell in cells.values():
                if cell.original and cell.has_commission and clean(cell.text) not in seen:
                    seen.append(clean(cell.text))
        return seen

    def search_text(self, audience: str = "customer") -> str:
        parts = []
        for _, cells in self.rows:
            for cell in cells.values():
                if cell.original:
                    parts.append(cell.text if audience == "agency" else cell.customer_text)
        return clean(" ".join(parts))


def _merged_ranges(archive: zipfile.ZipFile, path: str) -> list[tuple[int, int, int, int]]:
    try:
        xml = archive.read(path.lstrip("/")).decode("utf-8", "ignore")
    except KeyError:
        return []
    return [
        (column_index_from_string(c1), int(r1), column_index_from_string(c2), int(r2))
        for c1, r1, c2, r2 in _MERGE.findall(xml)
    ]


def parse_workbook(data: bytes) -> list[Sheet]:
    """يرجع الشيتات الظاهرة (غير المخفية) اللي بيها بيانات."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    archive = zipfile.ZipFile(io.BytesIO(data))
    sheets: list[Sheet] = []
    for ws in workbook.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        grid: dict[tuple[int, int], Cell] = {}
        for row in ws.iter_rows():
            for raw in row:
                cell = _make_cell(raw)
                if cell is not None:
                    grid[(raw.row, raw.column)] = cell
        if len(grid) < MIN_CELLS:
            continue
        count = len(grid)
        for c1, r1, c2, r2 in _merged_ranges(archive, getattr(ws, "_worksheet_path", "")):
            top = grid.get((r1, c1))
            if top is None:
                continue
            if c2 > c1:
                top.span_to = c2
            for r in range(r1 + 1, min(r2, r1 + 400) + 1):
                grid.setdefault((r, c1), Cell(top.text, top.number, top.unit, original=False, span_to=top.span_to))
        by_row: dict[int, dict[int, Cell]] = {}
        for (r, c), cell in grid.items():
            by_row.setdefault(r, {})[c] = cell
        rows = [(r, dict(sorted(by_row[r].items()))) for r in sorted(by_row)]
        sheets.append(Sheet(key=ws.title, title=clean(ws.title), rows=rows, cells=count))
    workbook.close()
    archive.close()
    return sheets


def export_url(sheet_id: str) -> str:
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"


def remove_commission_lines(text: str) -> str:
    """حماية أخيرة لنسخة الزبون: يشيل أي سطر يذكر العمولة من الناتج."""
    return "\n".join(line for line in text.split("\n") if not _COMMISSION.search(_INVISIBLE.sub("", line)))


def column_letter(index: int) -> str:
    return get_column_letter(index)

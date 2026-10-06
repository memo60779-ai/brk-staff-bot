"""فحص بدون إنترنت: شيت تجريبي + تليكرام وهمي.  التشغيل: python test_bot.py"""
import asyncio
import io
import os
from types import SimpleNamespace

import openpyxl

os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123456:TEST")
os.environ.setdefault("SHEET_ID", "TEST")
os.environ["ALLOWED_CHAT_IDS"] = "-100111"

import bot  # noqa: E402
import match  # noqa: E402
import pricing  # noqa: E402
import render  # noqa: E402


def sample_workbook() -> bytes:
    wb = openpyxl.Workbook()

    # شيت فنادق عادي: الفندق مدموج على مدتين، والبرنامج والتاريخ مدموجين على كل الصفوف
    ws = wb.active
    ws.title = "بيروت بغـــداد"
    ws.append(["البرنامج", "الموقع", "السنكل", "الرضيع", "الطفل بدون سرير", "المبيع", "عدد الايام", "تقييم الفندق", "اسم الفندق", "تاريخ السفر"])
    ws.append(["كروب بغداد بيروت ( عمولــه الشركات 20.000 دينار )"])
    ws.merge_cells("A2:J2")
    ws.append(["السعر يشمل\nتذكرة ذهاب وعودة", "الحمرا", 710000, 150000, 350000, 600000, "4 ايام / 3 ليالي", 3, "Villa Queens", "العراقية\n07/10/2026"])
    ws.append([None, None, 770000, 150000, 350000, 645000, "5 ايام / 4 ليالي"])
    ws.append([None, "الروشة", 795000, 150000, 350000, 660000, "4 ايام / 3 ليالي", 4, "City Suite"])
    ws.append([None, None, 870000, 150000, 350000, 700000, "5 ايام / 4 ليالي"])
    for ref in ("A3:A6", "J3:J6", "B3:B4", "H3:H4", "I3:I4", "B5:B6", "H5:H6", "I5:I6"):
        ws.merge_cells(ref)
    ws.append(["ملاحظة طويلة: يرجى التأكد من صلاحية الجواز ستة أشهر على الأقل قبل السفر، والشركة غير مسؤولة عن أي منع من السفر بسبب نقص الأوراق أو تأخر المسافر عن موعد الرحلة."])
    ws.merge_cells("A7:J7")
    for row in ws.iter_rows(min_row=3, max_row=6, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"

    # شيت بقسمين وعمود فندقين
    ws2 = wb.create_sheet("اسطنبول اور")
    ws2.append(["عمولة الشركات 10$"])
    ws2.merge_cells("A1:F1")
    ws2.append(["البرنامج", "الدبل", "السنكل", "فندق طرابزون", "فندق اسطنبول", "التقييم"])
    ws2.append(["طيران + فندق + استقبال", 450, 600, "SOGUT PARK", "PRENS HOTEL", 3])
    ws2.append([None, 520, 700, "CITY PORT", "PRENS HOTEL", 4])
    ws2.merge_cells("A3:A4")
    ws2.append(["برنامج VIP"])
    ws2.merge_cells("A5:F5")
    ws2.append([None, 900, 1200, "MOVENPICK", "HILTON", 5])
    for cell in ("B3", "C3", "B4", "C4", "B6", "C6"):
        ws2[cell].number_format = '"$"#,##0'

    # جدول عام بدون فنادق
    ws3 = wb.create_sheet("الفيزا")
    ws3.append(["الملاحظات", "السعر", "مدة الاصدار", "النوع", "الدول"])
    ws3.append(["فيزا الامارات"])
    ws3.merge_cells("A2:E2")
    ws3.append(["المبلغ غير مسترجع", 85, "3-5 ايام", "سياحية", "الامارات"])
    ws3.append([None, 20, "3-5 ايام", "طفل", None])
    ws3.merge_cells("A3:A4")
    ws3.merge_cells("E3:E4")

    old = wb.create_sheet("بيروت العيد قديم")
    for _ in range(4):
        old.append(["عرض قديم", 999999, "x", "y"])
    old.sheet_state = "hidden"
    wb.create_sheet("Sheet9")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class Chat:
    """محادثة وهمية تحفظ كل اللي رجعه البوت."""

    def __init__(self, chat_id=-100111):
        self.id = chat_id
        self.type = "group" if chat_id < 0 else "private"
        self.sent: list[tuple[str, object]] = []
        self.alerts: list[str] = []

    async def reply_text(self, text, reply_markup=None, **_):
        assert len(text) <= 4096, "رسالة أطول من حد تليكرام"
        self.sent.append((text, reply_markup))

    def message(self, text):
        return SimpleNamespace(text=text, reply_text=self.reply_text)

    def update(self, text):
        return SimpleNamespace(effective_chat=self, effective_user=SimpleNamespace(id=7), effective_message=self.message(text))

    def press(self, data):
        async def answer(text=None, **_):
            if text:
                self.alerts.append(text)

        query = SimpleNamespace(data=data, answer=answer, message=self.message(""))
        return SimpleNamespace(effective_chat=self, effective_user=SimpleNamespace(id=7), callback_query=query)

    def buttons(self, index=-1):
        return [button for row in self.sent[index][1].inline_keyboard for button in row]

    def texts(self):
        return [text for text, _ in self.sent]


async def run():
    sheets = pricing.parse_workbook(sample_workbook())
    assert [s.title for s in sheets] == ["بيروت بغداد", "اسطنبول اور", "الفيزا"], [s.title for s in sheets]
    beirut, istanbul, visa = sheets
    print("✓ قراءة الشيت: المخفي والفارغ متجاهلين والمدّات تنظفت من الأسماء")

    customer = render.render(beirut, "customer")
    agency = render.render(beirut, "agency")
    assert "عمول" not in customer and "عموله الشركات 20.000 دينار" in agency
    assert "كروب بغداد بيروت" in customer  # باقي الخلية بقى بعد قص العمولة
    assert "🔹 Villa Queens — الحمرا (3⭐)\n⬅️ 4 ايام / 3 ليالي | 600,000 د.ع | السنكل: 710,000 د.ع\n⬅️ 5 ايام / 4 ليالي | 645,000 د.ع | السنكل: 770,000 د.ع" in customer
    assert "🔹 City Suite — الروشة (4⭐)\n⬅️ 4 ايام / 3 ليالي | 660,000 د.ع" in customer
    assert customer.count("Villa Queens") == 1  # اسم الفندق مرة وحدة وتحته مدده
    assert "👶 الرضيع: 150,000 د.ع" in customer and customer.count("150,000") == 1  # القيمة الثابتة تنكتب مرة وحدة
    assert "✅ السعر يشمل\nتذكرة ذهاب وعودة" in customer and "🗓 العراقية\n07/10/2026" in customer
    assert "ملاحظة طويلة" in customer and "ملاحظة طويلة" not in render.render(beirut, "customer", brief=True)
    print("✓ قالب الفنادق: كل فندق وتحته مدده، العملة صحيحة، العمولة فقط بنسخة الشركات")

    text = render.render(istanbul, "customer")
    assert "عمول" not in text and "💼 عمولة الشركات 10$" in render.render(istanbul, "agency")
    assert "🔹 فندق اسطنبول: PRENS HOTEL\n➕ فندق طرابزون: SOGUT PARK (3⭐)\n⬅️ الدبل: 450$ | السنكل: 600$" in text
    assert "📍 برنامج VIP" in text and "➕ فندق طرابزون: MOVENPICK (5⭐)\n⬅️ الدبل: 900$ | السنكل: 1,200$" in text
    print("✓ شيت بقسمين وفندقين: كل سعر ويا فنادقه والقسم الثاني ورث عناوين الأعمدة")

    text = render.render(visa, "customer")
    assert "📍 فيزا الامارات" in text and "▫️ الدول: الامارات" in text
    assert "⬅️ 85$" in text and "⬅️ 20$" in text and "النوع: طفل" in text
    print("✓ جدول عام (فيزا): العناوين تتحول لحقول والسعر الصغير بدون رمز ينحسب دولار")

    for sheet in sheets:  # ولا رقم يضيع
        for audience in ("customer", "agency"):
            out = render.render(sheet, audience)
            for _, cells in sheet.rows:
                for cell in cells.values():
                    assert cell.number is None or cell.text in out, (sheet.title, cell.text)
    print("✓ كل رقم بالشيت موجود بالنص")

    bot.set_sheets(sheets)
    bot._cache["at"] = 10**12  # الشيت محمل: لا تحاول تتصل بالإنترنت
    context = SimpleNamespace(bot=SimpleNamespace(send_chat_action=lambda *a, **k: asyncio.sleep(0)))

    # 1) طلب واضح: عنوان بأزرار ثم النص نظيف
    chat = Chat()
    await bot.on_message(chat.update("دزلي نص كروبات بيروت"), context)
    assert chat.sent[0][0] == "📄 بيروت بغداد · نسخة الزبون (بدون عمولة)"
    assert [b.text for b in chat.buttons(0)] == ["💼 نسخة الشركات", "✂️ بدون الملاحظات"]
    body = chat.sent[1]
    assert body[1] is None and "عمول" not in body[0] and "07810105600" in body[0] and "600,000 د.ع" in body[0]
    print("✓ نسخة الزبون: بدون عمولة، الخاتمة مضافة، والنص بدون أزرار حتى ينسخ نظيف")

    # 2) زر نسخة الشركات
    await bot.on_button(chat.press(chat.buttons(0)[0].callback_data), context)
    assert "نسخة الشركات" in chat.sent[2][0] and "عموله الشركات 20.000 دينار" in chat.sent[3][0]
    assert chat.buttons(2)[0].text == "👤 نسخة الزبون"
    print("✓ زر نسخة الشركات يرجع النص ويا العمولة")

    # 3) كلمة شركات بالطلب
    chat = Chat()
    await bot.on_message(chat.update("نص اسطنبول للشركات"), context)
    assert "نسخة الشركات" in chat.sent[0][0] and "عمولة الشركات 10$" in chat.sent[1][0]
    print("✓ كلمة «شركات» بالطلب تطلع نسخة الشركات مباشرة")

    # 4) بحث باسم منطقة موجودة داخل الشيت
    chat = Chat()
    await bot.on_message(chat.update("الروشة"), context)
    assert chat.sent[0][0].startswith("📄 بيروت بغداد")
    print("✓ البحث باسم منطقة أو فندق يلكي الشيت")

    # 5) أكثر من عرض: أزرار اختيار وزر الكل
    two = pricing.parse_workbook(sample_workbook())
    two[1].title, two[1].key = "بيروت نجف", "بيروت نجف"
    bot.set_sheets(two)
    bot._cache["at"] = 10**12
    chat = Chat()
    await bot.on_message(chat.update("اريد عرض بيروت"), context)
    assert "لكيت 2 عروض" in chat.sent[0][0]
    assert [b.text for b in chat.buttons()] == ["بيروت بغداد", "بيروت نجف", "📚 الكل"]
    await bot.on_button(chat.press(chat.buttons()[-1].callback_data), context)
    labels = [t for t in chat.texts() if t.startswith("📄")]
    assert len(labels) == 2
    chat = Chat()
    await bot.on_message(chat.update("بيروت نجف"), context)
    assert chat.sent[0][0].startswith("📄 بيروت نجف")
    print("✓ طلب يطابق أكثر من شيت: أزرار اختيار وزر «الكل»، والاسم الكامل يختار واحد")

    # 6) كلام عادي بالكروب: البوت يسكت. طلب لعرض ما موجود: يرد بالقائمة
    chat = Chat()
    await bot.on_message(chat.update("صباح الخير شباب"), context)
    await bot.on_message(chat.update("منو راح يكمل ملف زبون بيروت اللي اجه البارحة العصر"), context)
    assert chat.sent == []
    await bot.on_message(chat.update("اريد عرض لندن"), context)
    assert "ما لكيت" in chat.sent[0][0] and len(chat.buttons()) == 3
    await bot.on_message(chat.update("شنو العروض"), context)
    assert "العروض الفعالة" in chat.sent[1][0]
    print("✓ السوالف العادية ما يرد عليها، والطلب الغلط يرجع قائمة العروض")

    # 7) محادثة خارج كروب الموظفين
    outsider = Chat(chat_id=999)
    await bot.on_message(outsider.update("دزلي نص بيروت"), context)
    await bot.on_button(outsider.press(f"s|a|{match.sheet_id(two[0])}"), context)
    assert outsider.sent == [] and outsider.alerts == ["هذه المحادثة غير مضافة للبوت"]
    print("✓ خارج كروب الموظفين: البوت ما يرد ولا يشتغل زر الشركات")

    # 8) زر قديم بعد ما انحذف العرض من الشيت
    chat = Chat()
    await bot.on_button(chat.press("s|c|deadbeef"), context)
    assert chat.sent == [] and "/sheets" in chat.alerts[0]
    print("✓ زر لعرض انحذف من الشيت: تنبيه بدل خطأ")

    # 9) نص طويل ينقسم على حد تليكرام بدون ما يقطع سطر
    long_text = "\n".join(f"سطر {i} " + "س" * 80 for i in range(120))
    chunks = render.split_message(long_text)
    assert len(chunks) > 1 and all(len(c) <= 3900 for c in chunks) and "\n".join(chunks) == long_text
    print("✓ النص الطويل ينقسم لرسائل بدون ما يضيع شي")

    # 10) قراءة معرف الكروب بكل أشكاله
    assert bot.parse_chat_ids("-5492479479") == {-5492479479}
    assert bot.parse_chat_ids("5492479479-") == {-5492479479}  # الناقص بعد الرقم مثل ما يطلع بالعربي
    assert bot.parse_chat_ids(" -100123 , -100456 ") == {-100123, -100456}
    bot.ALLOWED = bot.parse_chat_ids("5492479479")  # انكتب بدون ناقص
    assert bot.chat_allowed(SimpleNamespace(id=-5492479479, type="group"))
    bot.ALLOWED = {-5492479479}
    assert not bot.chat_allowed(SimpleNamespace(id=5492479479, type="private"))  # محادثة خاصة بنفس الرقم ما تنقبل
    assert not bot.chat_allowed(SimpleNamespace(id=-777, type="group"))
    bot.ALLOWED = {-100111}
    print("✓ معرف الكروب ينقبل حتى لو الناقص بعد الرقم أو مفقود، والمحادثات الخاصة ما تنقبل")

    from telegram.ext import Application

    Application.builder().token(os.environ["TELEGRAM_BOT_TOKEN"]).concurrent_updates(True).build()
    print("✓ تطبيق تليكرام ينبني بدون أخطاء")


if __name__ == "__main__":
    asyncio.run(run())
    print("\nكل الفحوصات نجحت")
